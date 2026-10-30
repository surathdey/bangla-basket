"""BongConnekt grocery catalogue template — build + parse.

The Admin panel offers a two-step self-serve catalog flow:
  1. Download Template  -> `build_template_xlsx()` returns a ready-to-fill .xlsx
                            with the required column headers + a few sample rows.
  2. Upload Template    -> `parse_catalogue_xlsx()` reads a filled workbook and
                            returns grouped products (rows sharing the same
                            "Item" name collapse into one product with variants).

Images are intentionally NOT part of this flow — the admin adds photos later
from the item editor. New products come in with `image_url == ""`.
"""
from __future__ import annotations

from io import BytesIO
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# Column order MUST match the customer-supplied template exactly.
TEMPLATE_HEADERS = [
    "Primary Category",
    "All Categories",
    "Item",
    "Variant",
    "Current Price (USD)",
    "Compare-at Price (USD)",
    "Savings (USD)",
    "Discount",
    "On Sale",
    "Available",
    "Delivery Scope",
    "Weight (g)",
    "Description",
]

# Column indexes (0-based) for readable parsing.
C_PRIMARY_CAT = 0
C_ALL_CATS = 1
C_ITEM = 2
C_VARIANT = 3
C_PRICE = 4
C_COMPARE = 5
C_SAVINGS = 6
C_DISCOUNT = 7
C_ON_SALE = 8
C_AVAILABLE = 9
C_DELIVERY = 10
C_WEIGHT = 11
C_DESC = 12

# A handful of sample rows so the shop owner sees exactly how to fill it in —
# including one product ("Whole Chicken") that has two variants.
SAMPLE_ROWS = [
    ["Rice & Grains", "Rice & Grains", "Premium Basmati Rice 5kg", "Standard",
     18.99, 22.99, 4.00, 0.17, "Yes", "Yes", "GTA", 5000,
     "Aged long-grain basmati rice — fluffy, aromatic, perfect for biryani."],
    ["Chicken", "Chicken, Halal Meat", "Whole Halal Chicken", "Small (approx 2.5 lb)",
     12.99, None, 0, 0, "No", "Yes", "GTA", 1134,
     "Fresh halal whole chicken, cleaned and cut to your preference."],
    ["Chicken", "Chicken, Halal Meat", "Whole Halal Chicken", "Large (approx 4 lb)",
     18.99, None, 0, 0, "No", "Yes", "GTA", 1814,
     "Fresh halal whole chicken, cleaned and cut to your preference."],
    ["Fish", "Fish, Fresh Seafood", "Rohu Fish (Whole)", "Standard",
     14.99, 16.99, 2.00, 0.12, "Yes", "Yes", "GTA", 1500,
     "Fresh water rohu — cleaned, gutted and cut into curry pieces."],
]


def build_template_xlsx() -> bytes:
    """Return the .xlsx template (headers + sample rows) as raw bytes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="011838")
    center = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.append(TEMPLATE_HEADERS)
    for col_idx, _ in enumerate(TEMPLATE_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for row in SAMPLE_ROWS:
        ws.append(row)

    # Reasonable column widths for readability.
    widths = [16, 20, 30, 30, 16, 18, 12, 10, 9, 10, 14, 11, 42]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _to_cents(v) -> int:
    try:
        if v is None or str(v).strip() == "":
            return 0
        return int(round(float(v) * 100))
    except (ValueError, TypeError):
        return 0


def _price_cents_checked(v) -> tuple[int, bool]:
    """Parse a price cell. Returns (cents, ok). ok is False when the value is
    missing, non-numeric, zero or negative — so invalid rows can be rejected
    instead of being silently imported as $0.00 items."""
    if v is None or str(v).strip() == "":
        return 0, False
    try:
        cents = int(round(float(v) * 100))
    except (ValueError, TypeError):
        return 0, False
    if cents <= 0:
        return 0, False
    return cents, True


def _to_weight(v) -> int:
    try:
        if v is None or str(v).strip() == "":
            return 0
        return int(round(float(v)))
    except (ValueError, TypeError):
        return 0


def _clean_desc(v) -> str:
    s = str(v or "").strip()
    # Ignore the placeholder shipped in the sample template.
    if not s or s.lower().startswith("<auto"):
        return ""
    return s


def parse_catalogue_xlsx(content: bytes) -> tuple[list[dict], list[dict]]:
    """Parse an uploaded catalogue workbook into grouped product dicts.

    Rows are grouped by the (trimmed) "Item" name. Each group becomes one
    product; every row in the group becomes a variant. Returns a tuple of
    (products, errors) where each error is {row, reason} (row = 1-based sheet
    row number). Invalid rows (missing item name, or missing / non-numeric /
    zero / negative price) are rejected and reported — never imported as
    $0.00 items.
    """
    wb = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []

    # Skip the header row (detected if first cell equals the known header).
    start = 1 if str((rows[0] or [None])[0] or "").strip().lower() == "primary category" else 0

    errors: list[dict] = []
    # group -> list of (variant_dict)
    groups: "OrderedDict[str, list]" = OrderedDict()
    meta: "OrderedDict[str, dict]" = OrderedDict()
    for i in range(start, len(rows)):
        r = rows[i]
        sheet_row = i + 1  # 1-based row number as seen in Excel
        if r is None:
            continue
        # pad short rows
        r = list(r) + [None] * (len(TEMPLATE_HEADERS) - len(r))
        # Skip fully blank rows silently.
        if not any(str(c or "").strip() for c in r):
            continue
        title = str(r[C_ITEM] or "").strip()
        if not title:
            errors.append({"row": sheet_row, "reason": "Missing item name"})
            continue
        price_cents, price_ok = _price_cents_checked(r[C_PRICE])
        if not price_ok:
            errors.append({"row": sheet_row, "reason": f"Invalid or missing price for '{title}'"})
            continue

        vtitle = str(r[C_VARIANT] or "Standard").strip() or "Standard"
        compare = _to_cents(r[C_COMPARE])
        avail = str(r[C_AVAILABLE] or "").strip().lower() != "no"
        on_sale_row = str(r[C_ON_SALE] or "").strip().lower() == "yes"
        desc = _clean_desc(r[C_DESC])

        groups.setdefault(title, []).append({
            "title": vtitle,
            "price_cents": price_cents,
            "compare_at_cents": compare,
            "available": avail,
            "weight_g": _to_weight(r[C_WEIGHT]),
        })
        m = meta.setdefault(title, {"category": "", "all_cats": "", "description": "", "on_sale": False})
        if not m["category"]:
            m["category"] = str(r[C_PRIMARY_CAT] or "").strip()
        if not m["all_cats"]:
            m["all_cats"] = str(r[C_ALL_CATS] or "").strip()
        if not m["description"] and desc:
            m["description"] = desc
        if on_sale_row:
            m["on_sale"] = True

    products: list[dict] = []
    for title, variants in groups.items():
        m = meta[title]
        category = m["category"] or "Grocery & Pantry"
        has_variants = len(variants) > 1 or (variants and variants[0]["title"].lower() != "standard")
        avail_priced = [v for v in variants if v["available"]]
        pick = min(avail_priced or variants, key=lambda v: v["price_cents"])

        products.append({
            "title": title,
            "category": category,
            "all_categories": m["all_cats"],
            "description": m["description"],
            "on_sale": m["on_sale"],
            "price_cents": pick["price_cents"],
            "compare_at_cents": pick["compare_at_cents"],
            "has_variants": bool(has_variants),
            "in_stock": any(v["available"] for v in variants),
            "variants": variants,
        })

    return products, errors
